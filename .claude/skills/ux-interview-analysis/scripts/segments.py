#!/usr/bin/env python3
"""Transcript runs -> ux_transcript rows + readable transcripts (Markdown, .docx, plain text).

Pure: reads the ux_runs export (JSON from the "UX Harness: Rows" read), writes files.
No network, no credentials. All labels come from locales/<lang>.yaml.

Usage:
    python scripts/segments.py runs.json --interview S1-U01 --lang cs --duration 2582.9 \
        --rows-out rows.json --md-out t.md [--docx-out t.docx] [--txt-out t.txt] \
        [--timestamps absolute|relative]

- Keeps the newest successful transcript run per clip window (window = videoMetadata.startOffset).
- Accepts transcript.v1 (text_cs) and transcript.v2 (text + language) outputs.
- --timestamps relative: the model reported clip-relative times; the window start is added.
- QA report (stderr): per-window counts, speakers, detected languages and every rule violation.
  Nothing is "fixed" silently: out-of-window / non-monotonic segments are kept and flagged.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

import yaml

SKILL = Path(__file__).resolve().parent.parent


def load_locale(code: str) -> dict:
    path = SKILL / "locales" / f"{code}.yaml"
    if not path.exists():
        sys.exit(f"no locale {path} - add one (see locales/README.md)")
    return yaml.safe_load(path.read_text(encoding="utf-8"))


def to_seconds(ts: str) -> float:
    parts = [float(p) for p in str(ts).strip().split(":")]
    if not 1 <= len(parts) <= 3:
        raise ValueError(f"bad timestamp {ts!r}")
    total = 0.0
    for p in parts:
        total = total * 60 + p
    return total


def mmss(seconds: float) -> str:
    s = int(round(seconds))
    return f"{s // 60:02d}:{s % 60:02d}"


def window_of(run: dict) -> tuple[float, float | None]:
    params = json.loads(run.get("params_json") or "{}")
    for content in params.get("contents", []):
        for part in content.get("parts", []):
            clip = part.get("videoMetadata")
            if clip:
                start = float(str(clip.get("startOffset", "0s")).rstrip("s") or 0)
                end = clip.get("endOffset")
                return start, (float(str(end).rstrip("s")) if end else None)
    return 0.0, None


def exec_no(run_id: str) -> int:
    m = re.search(r"-(\d+)$", run_id or "")
    return int(m.group(1)) if m else -1


def header_lines(loc: dict, interview: str, duration: float | None) -> list[str]:
    t = loc["transcript"]
    intro = list(t["intro"])
    if duration:
        intro[1] += t["duration"].format(duration=mmss(duration))
    return intro


def write_docx(path: Path, loc: dict, interview: str, rows: list, source: str, notes: list,
               duration: float | None) -> None:
    from docx import Document  # lazy: only needed for the client document
    from docx.shared import Pt

    t = loc["transcript"]
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(t["heading"].format(id=interview), level=1)
    for line in header_lines(loc, interview, duration):
        doc.add_paragraph(line)
    if notes:
        doc.add_paragraph(f"{t['audio_notes']}: " + " | ".join(notes))
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, t["table_header"]):
        cell.text = label
        cell.paragraphs[0].runs[0].bold = True
    for r in rows:
        cells = table.add_row().cells
        cells[0].text = f"{mmss(r['t_start_s'])}–{mmss(r['t_end_s'])}"
        cells[1].text = loc["roles"].get(r["speaker"], r["speaker"])
        cells[2].text = r["text_cs"]
    doc.add_paragraph()
    doc.add_paragraph(source).runs[0].font.size = Pt(8)
    doc.save(path)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("runs", type=Path, help="Rows read output for ux_runs")
    ap.add_argument("--interview", required=True)
    ap.add_argument("--lang", required=True, help="documentation language (locales/<lang>.yaml)")
    ap.add_argument("--timestamps", choices=["absolute", "relative"], default="absolute")
    ap.add_argument("--duration", type=float, help="recording length in seconds (for the header)")
    ap.add_argument("--rows-out", type=Path, required=True)
    ap.add_argument("--md-out", type=Path, required=True)
    ap.add_argument("--docx-out", type=Path, help="client-ready transcript (needs python-docx)")
    ap.add_argument("--txt-out", type=Path, help="plain-text transcript for the Publish Doc step")
    a = ap.parse_args()
    loc = load_locale(a.lang)
    roles = loc["roles"]

    data = json.loads(a.runs.read_text(encoding="utf-8"))
    runs = [r for r in data.get("rows", [])
            if r.get("interview_id") == a.interview and r.get("stage") == "transcript" and r.get("status") == "ok"]
    newest: dict[float, dict] = {}
    for r in runs:
        start, _ = window_of(r)
        if start not in newest or exec_no(r["run_id"]) > exec_no(newest[start]["run_id"]):
            newest[start] = r
    if not newest:
        sys.exit("no successful transcript runs for " + a.interview)

    segments, notes, problems, languages = [], [], [], {}
    for start in sorted(newest):
        run = newest[start]
        _, end = window_of(run)
        out = json.loads(run["output_json"])
        offset = start if a.timestamps == "relative" else 0.0
        win = f"{mmss(start)}–{mmss(end) if end else 'end'}"
        if out.get("notes"):
            notes.append(f"{win}: {out['notes']}")
        if out.get("language"):
            languages[out["language"]] = languages.get(out["language"], 0) + 1
        prev = None
        counts: dict[str, int] = {}
        for i, seg in enumerate(out.get("segments", [])):
            try:
                t0 = to_seconds(seg["t_start"]) + offset
                t1 = to_seconds(seg["t_end"]) + offset
            except (KeyError, ValueError) as e:
                problems.append(f"{win} #{i}: unreadable time ({e})")
                continue
            text = (seg.get("text") or seg.get("text_cs") or "").strip()
            role = seg.get("speaker", "unclear")
            if not text:
                problems.append(f"{win} #{i}: empty text")
            if t1 < t0:
                problems.append(f"{win} #{i} {mmss(t0)}: ends before it starts")
            if t0 < start - 5 or (end is not None and t0 > end + 5):
                problems.append(f"{win} #{i} {mmss(t0)}: outside its clip window")
            if prev is not None and t0 + 1 < prev:
                problems.append(f"{win} #{i} {mmss(t0)}: goes back in time (previous {mmss(prev)})")
            prev = t0
            counts[role] = counts.get(role, 0) + 1
            segments.append({"t0": t0, "t1": t1, "speaker": role, "text": text, "run_id": run["run_id"]})
        print(f"{win}: {sum(counts.values())} segments {counts} lang={out.get('language', '?')} "
              f"run={run['run_id']} model={run.get('model_version') or run.get('model')}", file=sys.stderr)

    segments.sort(key=lambda s: s["t0"])
    # Column name text_cs is historical; it holds the text in the language actually spoken.
    rows = [{
        "segment_id": f"{a.interview}-T{n:04d}",
        "interview_id": a.interview,
        "run_id": s["run_id"],
        "t_start_s": round(s["t0"], 2),
        "t_end_s": round(s["t1"], 2),
        "speaker": s["speaker"],
        "text_cs": s["text"],
    } for n, s in enumerate(segments, 1)]
    a.rows_out.write_text(json.dumps({"op": "write", "table": "ux_transcript", "rows": rows},
                                     ensure_ascii=False, indent=1), encoding="utf-8")

    first = next(iter(newest.values()))
    t = loc["transcript"]
    source = t["source"].format(model=first.get("model"), model_version=first.get("model_version"),
                                prompt=first.get("prompt_version"),
                                runs=", ".join(r["run_id"] for r in newest.values()))
    head = header_lines(loc, a.interview, a.duration)

    md = [f"# {t['heading'].format(id=a.interview)}", ""] + head + [""]
    if notes:
        md += [f"{t['audio_notes']}:"] + [f"- {n}" for n in notes] + [""]
    md += ["---", ""]
    for r in rows:
        md += [f"**[{mmss(r['t_start_s'])}–{mmss(r['t_end_s'])}] {roles.get(r['speaker'], r['speaker'])}:** {r['text_cs']}", ""]
    md += ["---", source]
    a.md_out.write_text("\n".join(md), encoding="utf-8")

    if a.docx_out:
        write_docx(a.docx_out, loc, a.interview, rows, source, notes, a.duration)

    if a.txt_out:  # plain text -> published as a Google Doc by the n8n Publish step
        txt = [t["heading"].format(id=a.interview), ""] + head + [""]
        txt += [f"[{mmss(r['t_start_s'])}–{mmss(r['t_end_s'])}] {roles.get(r['speaker'], r['speaker'])}: {r['text_cs']}"
                for r in rows]
        txt += ["", source]
        a.txt_out.write_text("\n".join(txt) + "\n", encoding="utf-8")

    if languages:
        print(f"detected languages per window: {languages}", file=sys.stderr)
        if any(code != a.lang for code in languages):
            problems.append(f"spoken language {sorted(languages)} differs from --lang {a.lang}: check doc_language.py")
    print(f"TOTAL {len(rows)} segments; {len(problems)} QA flags", file=sys.stderr)
    for p in problems:
        print("  QA: " + p, file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
